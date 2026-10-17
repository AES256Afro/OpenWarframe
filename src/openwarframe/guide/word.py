"""Write a MasteryPlan as an editable Word document (mastery-plan.docx).

Same sections as the HTML page, in print order. Uses only built-in Word styles plus
direct formatting, so it opens cleanly in Word, Google Docs and LibreOffice.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .mastery import MasteryPlan

GOLD = RGBColor(0x8F, 0x6A, 0x17)
CYAN = RGBColor(0x0D, 0x78, 0x91)
INK = RGBColor(0x17, 0x20, 0x33)
MUTED = RGBColor(0x56, 0x61, 0x78)
HEADER_FILL = "F3E7C9"
DISPLAY = "Georgia"
BODY = "Calibri"


def _shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _style_document(doc) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = BODY
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY)
    normal.paragraph_format.space_after = Pt(4)
    for level, size, color in ((1, 17, GOLD), (2, 13, INK), (3, 11, INK)):
        st = doc.styles[f"Heading {level}"]
        st.font.name = DISPLAY
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = color
        st.element.rPr.rFonts.set(qn("w:eastAsia"), DISPLAY)
        st.paragraph_format.space_before = Pt(14 if level == 1 else 10)
        st.paragraph_format.space_after = Pt(4)
        st.paragraph_format.keep_with_next = True
    title = doc.styles["Title"]
    title.font.name = DISPLAY
    title.font.size = Pt(28)
    title.font.color.rgb = INK
    title.element.rPr.rFonts.set(qn("w:eastAsia"), DISPLAY)
    for section in doc.sections:
        section.left_margin = section.right_margin = Cm(2)
        section.top_margin = section.bottom_margin = Cm(1.8)


def _label(paragraph, text: str, color: RGBColor = GOLD) -> None:
    run = paragraph.add_run(text.upper() + "  ")
    run.bold = True
    run.font.size = Pt(8.5)
    run.font.color.rgb = color


def _muted(doc, text: str) -> None:
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.color.rgb = MUTED
    r.italic = True


def _table(doc, header: list[str], rows: list[list[str]], *, right: set[int] | None = None, widths=None):
    right = right or set()
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        cell = t.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(h.upper())
        run.bold = True
        run.font.size = Pt(8.5)
        run.font.color.rgb = MUTED
        _shade(cell, HEADER_FILL)
        if i in right:
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = v
            if i in right:
                cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return t


def build_mastery_docx(plan: MasteryPlan):
    doc = Document()
    _style_document(doc)
    who = plan.display_name.rstrip("") if plan.display_name else "Your account"

    # title block
    eyebrow = doc.add_paragraph()
    _label(eyebrow, "Mastery plan · built from your profile")
    doc.add_paragraph(f"Road to Mastery {plan.target_rank}", style="Title")
    meta = doc.add_paragraph()
    meta.add_run(
        f"{who} · {plan.platform.upper()} · MR {plan.rank} · snapshot {plan.snapshot_taken_at:%d %b %Y}"
    ).font.color.rgb = MUTED
    n = len(plan.tests)
    doc.add_paragraph(
        f"{n} rank{'s' if n != 1 else ''} to go and about {plan.needed_to_target:,} mastery. Follow the path in "
        "order, or pick by mood when you want something different."
    )

    # where you stand
    doc.add_heading("Where you stand", level=1)
    p = doc.add_paragraph()
    p.add_run("Estimated mastery ").bold = False
    p.add_run(f"≈{plan.estimate:,}").bold = True
    if plan.ranks:
        p.add_run(f"   ·   to MR {plan.ranks[0].rank} ")
        p.add_run(f"≈{plan.ranks[0].to_earn:,}").bold = True
    p.add_run(f"   ·   to MR {plan.target_rank} ")
    p.add_run(f"≈{plan.needed_to_target:,}").bold = True
    _muted(
        doc,
        "The estimate adds item XP, intrinsics, star chart, junctions and Steel Path nodes, and may be off by a "
        "few thousand. Your in-game Profile shows the exact bar.",
    )
    if plan.ranks:
        _table(
            doc,
            ["Rank", "Title", "Total needed", "To earn", "Test to pass"],
            [[str(r.rank), r.title, f"{r.threshold:,}", f"{r.to_earn:,}", r.test] for r in plan.ranks],
            right={2, 3},
            widths=[1.4, 3.4, 2.8, 2.4, 6.0],
        )

    # rules
    doc.add_heading("Rules that make it faster", level=1)
    for t in plan.tips:
        para = doc.add_paragraph(style="List Bullet")
        para.add_run(t.title + " ").bold = True
        para.add_run(t.text)

    # tonight
    doc.add_heading("Do these first, tonight", level=1)
    for item in plan.tonight:
        para = doc.add_paragraph()
        para.paragraph_format.left_indent = Cm(0.2)
        para.add_run("☐  ")
        para.add_run(item.title).bold = True
        gain = para.add_run(f"   {item.gain_label}")
        gain.bold = True
        gain.font.color.rgb = GOLD
        detail = doc.add_paragraph(item.detail)
        detail.paragraph_format.left_indent = Cm(0.9)
        for r in detail.runs:
            r.font.color.rgb = MUTED
            r.font.size = Pt(9.5)

    # path
    doc.add_heading("The path, one rank at a time", level=1)
    for ph in plan.phases:
        need = "already there" if ph.already_there else f"≈{ph.need:,}"
        doc.add_heading(f"{ph.from_rank} → {ph.to_rank}   {ph.title}", level=2)
        head = doc.add_paragraph()
        _label(head, f"Need {need}", MUTED)
        _label(head, f"Test: {ph.test}", CYAN)
        if ph.already_there:
            doc.add_paragraph(
                "Your estimate already covers this rank. Take the test as soon as the previous one unlocks."
            )
        else:
            if ph.blurb:
                _muted(doc, ph.blurb)
            for s in ph.steps:
                para = doc.add_paragraph(style="List Number")
                para.add_run(s.title).bold = True
                m = para.add_run(f"  +{s.mastery:,}")
                m.bold = True
                m.font.color.rgb = GOLD
                para.add_run("\n" + ", ".join(s.items) + ".")
            if ph.planned < ph.need:
                _muted(
                    doc,
                    f"The sources this plan knows about add {ph.planned:,} of the {ph.need:,} this rank needs. "
                    "Trade or event items can cover the rest.",
                )
        for label, text, color in (
            ("If you get tired", ph.tired, CYAN),
            ("If it is too hard", ph.too_hard, RGBColor(0xA8, 0x65, 0x0C)),
            ("When you are there", ph.done, RGBColor(0x1F, 0x8A, 0x55)),
        ):
            para = doc.add_paragraph()
            para.paragraph_format.left_indent = Cm(0.4)
            _label(para, label, color)
            para.add_run(text)

    # moods
    doc.add_heading("Pick by mood", level=1)
    _muted(doc, "Every option earns affinity or mastery.")
    for c in plan.mood_cards:
        doc.add_heading(c.name, level=3)
        tag = doc.add_paragraph()
        _label(tag, " · ".join(c.mood_labels))
        for label, text in (
            ("Where", c.where),
            ("Levels", c.levels),
            ("Too hard", c.too_hard),
            ("Bored", c.tired),
        ):
            para = doc.add_paragraph()
            para.paragraph_format.left_indent = Cm(0.4)
            _label(para, label, MUTED)
            para.add_run(text)

    # budget
    doc.add_heading(f"Where {plan.needed_to_target:,} mastery comes from", level=1)
    _muted(
        doc,
        f"You need about {plan.needed_to_target:,}. Everything except trading adds up to about "
        f"{plan.budget_without_trade:,}.",
    )
    _table(
        doc,
        ["Source", "Items", "Mastery"],
        [[b.label + (" (trade)" if b.trade else ""), str(b.count), f"{b.mastery:,}"] for b in plan.budget],
        right={1, 2},
        widths=[10.0, 2.0, 3.6],
    )

    # tests
    if plan.tests:
        doc.add_heading("The tests", level=1)
        _muted(doc, " ".join(plan.test_rules))
        for t in plan.tests:
            doc.add_heading(f"MR {t.rank} · {t.title}: {t.name}", level=2)
            doc.add_paragraph(t.objective)
            for tip in t.tips:
                doc.add_paragraph(tip, style="List Bullet")
            para = doc.add_paragraph()
            _label(para, "Bring")
            para.add_run(t.bring)

    # footer
    doc.add_heading("Sources and notes", level=1)
    _muted(
        doc,
        f"Built by OpenWarframe on {plan.generated_at:%d %b %Y} from the snapshot taken "
        f"{plan.snapshot_taken_at:%d %b %Y %H:%M} UTC.",
    )
    for s in plan.sources:
        doc.add_paragraph(f"{s['title']}: {s['url']}", style="List Bullet")
    parts = ", ".join(f"{k.replace('_', ' ')} {v:,}" for k, v in plan.estimate_parts.items())
    _muted(doc, f"Estimate: {parts}. Total ≈{plan.estimate:,}.")
    for n_ in plan.notes:
        _muted(doc, n_)
    return doc


def write_mastery_docx(plan: MasteryPlan, path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    build_mastery_docx(plan).save(str(path))
    return path
