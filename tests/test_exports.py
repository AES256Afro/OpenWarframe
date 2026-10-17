"""Exports from a tiny codex and a synthetic profile. No network."""

from __future__ import annotations

import csv
import json
import sqlite3
from datetime import UTC, datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook

from openwarframe.account import parse_profile
from openwarframe.codex import Codex, CodexPayloads, build_codex
from openwarframe.exports import run_exports
from openwarframe.exports.brief import build_brief
from openwarframe.report import Report, build_report

from .test_codex_build import RHINO, _drops, _pe, _plus

BRATON = "/Lotus/Weapons/Tenno/Rifle/Braton"


def _codex(tmp_path: Path) -> Codex:
    pe = _pe()
    pe.manifests["ExportWeapons"]["ExportWeapons"] = [
        {"uniqueName": BRATON, "name": "Braton", "productCategory": "LongGuns", "masteryReq": 0},
        {
            "uniqueName": "/Lotus/Weapons/Tenno/Rifle/Soma",
            "name": "Soma",
            "productCategory": "LongGuns",
            "masteryReq": 6,
        },
        {
            "uniqueName": "/Lotus/Weapons/Hidden/Part",
            "name": "Hidden Part",
            "productCategory": "LongGuns",
            "excludeFromCodex": True,
        },
    ]
    path = tmp_path / "codex.sqlite"
    conn = sqlite3.connect(path)
    build_codex(conn, CodexPayloads(public_export=pe, plus=_plus(), drops=_drops(), market=[]))
    conn.close()
    return Codex(path)


def _profile() -> dict:
    return {
        "Results": [
            {
                "DisplayName": "Tester",
                "PlayerLevel": 3,
                "GuildName": "Clan",
                "LoadOutInventory": {
                    "XPInfo": [{"ItemType": BRATON, "XP": 450_000}, {"ItemType": RHINO, "XP": 250_000}]
                },
                "Missions": [
                    {"Tag": "SolNode94", "Completes": 2, "Tier": 1},
                    {"Tag": "EarthToMarsJunction", "Completes": 1, "Tier": 1},
                    {"Tag": "MarsToCeresJunction", "Completes": 1},
                ],
                "Affiliations": [{"Tag": "ArbitersSyndicate", "Standing": 1000, "Title": 5}],
                "PlayerSkills": {"LPS_GUNNERY": 2},
                "ChallengeProgress": [{"Name": "X", "Progress": 4}, {"Name": "BladeMastery1", "Progress": 2}],
            }
        ],
        "Stats": {
            "TimePlayedSec": 3600,
            "Weapons": [{"type": BRATON, "kills": 10, "headshots": 2, "equipTime": 7200, "xp": 450000}],
        },
    }


@pytest.fixture
def report(tmp_path: Path) -> Report:
    cx = _codex(tmp_path)
    snap = parse_profile(_profile(), codex=cx, taken_at=datetime(2026, 10, 16, tzinfo=UTC))
    return build_report(snap, cx)


def test_report_contents(report: Report):
    h = report.headline
    assert h.mastery_rank == 3 and h.items_mastered == 1 and h.items_leveling == 1
    names = {m.name for m in report.missing}
    assert "Soma" in names and "Hidden Part" not in names and "Braton" not in names
    soma = next(m for m in report.missing if m.name == "Soma")
    assert soma.above_your_mr and soma.wiki_url.endswith("/Soma")
    rhino = next(o for o in report.owned if o.unique_name == RHINO)
    assert not rhino.mastered and rhino.xp_to_30 == 900_000 - 250_000
    ms = report.mastery
    assert ms.from_intrinsics == 3000
    assert ms.from_items == 3000 + 3000  # Braton rank 30, Rhino rank 15
    assert ms.visible_total == ms.from_items + ms.from_intrinsics
    assert ms.from_starchart == 60 and ms.from_steel_path == 60  # one node, about 60 each
    assert ms.from_junctions == 3000  # two junctions + one Steel Path junction
    assert ms.estimated_total == ms.visible_total + 60 + 60 + 3000
    assert report.nodes[0].completed and report.nodes[0].steel_path
    assert {a.name for a in report.achievements} == {"Archivist", "Blade Mastery1"}
    assert report.usage[0].equip_hours == 2.0
    assert report.notes  # profile-only note


def test_all_formats(report: Report, tmp_path: Path):
    out = tmp_path / "out"
    results = run_exports(report, out, ["json", "sqlite", "csv", "xlsx", "llm"])
    assert [r.format for r in results] == ["json", "sqlite", "csv", "xlsx", "llm"]

    data = json.loads((out / "account.json").read_text(encoding="utf-8"))
    assert data["schema_version"] == "openwarframe.account/1" and data["headline"]["mastery_rank"] == 3

    conn = sqlite3.connect(out / "account.sqlite")
    tables = {r[0] for r in conn.execute("SELECT name FROM sqlite_master WHERE type='table'")}
    assert {"meta", "summary", "missing_items", "owned_items", "star_chart", "syndicates"} <= tables
    assert conn.execute("SELECT COUNT(*) FROM owned_items").fetchone()[0] == 2
    conn.close()

    with (out / "csv" / "missing_items.csv").open(encoding="utf-8-sig") as fh:
        rows = list(csv.DictReader(fh))
    assert any(r["name"] == "Soma" for r in rows)

    xlsx = next(out.glob("*.xlsx"))
    wb = load_workbook(xlsx)
    assert wb.sheetnames[0] == "Summary" and "Never leveled" in wb.sheetnames
    assert wb["Summary"]["B2"].value.startswith("Tester")
    ws = wb["Never leveled"]
    assert ws["A1"].value == "Name" and ws.auto_filter.ref

    md = (out / "account-brief.md").read_text(encoding="utf-8")
    assert md.startswith("---\nschema: openwarframe.brief/1") and "# Mastery by category" in md
    assert f"Estimated total mastery XP (ESTIMATE) about {report.mastery.estimated_total:,}" in md
    brief = json.loads((out / "account-brief.json").read_text(encoding="utf-8"))
    assert brief["headline"]["mastery_rank"] == 3


def test_share_mode_strips_identity(tmp_path: Path):
    cx = _codex(tmp_path)
    snap = parse_profile(_profile(), codex=cx, taken_at=datetime(2026, 10, 16, tzinfo=UTC))
    r = build_report(snap, cx, anonymize=True)
    assert r.identity.display_name is None and r.identity.clan is None and r.identity.anonymized
    text, _ = build_brief(r)
    assert "Tester" not in text and "Clan" not in text


def test_schema_file_is_current():
    path = Path(__file__).resolve().parents[1] / "docs" / "schema" / "account.schema.json"
    committed = json.loads(path.read_text(encoding="utf-8"))
    assert committed["properties"] == Report.model_json_schema()["properties"], "run: openwarframe schema"


def test_syndicate_labels():
    from openwarframe.report.build import syndicate_label

    assert (
        syndicate_label("Nightwave", "RadioLegionIntermission16Syndicate", True)
        == "Nightwave Intermission 16"
    )
    assert syndicate_label("Nightwave", "RadioLegionSyndicate", True) == "Nightwave Series 1"
    assert syndicate_label("Nightwave", "RadioLegion3Syndicate", True) == "Nightwave Series 3"
    assert syndicate_label("Ostron", "CetusSyndicate", False) == "Ostron"
    assert syndicate_label("Same", "OtherSyndicate", True) == "Same (Other)"


def test_starchart_mastery_estimate():
    from openwarframe.report.build import starchart_mastery

    assert starchart_mastery(0) == 0
    assert starchart_mastery(241) == 14_569
    assert starchart_mastery(300) == 14_569  # capped at the mastery nodes
