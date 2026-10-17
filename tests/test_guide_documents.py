"""PDF, Word, toolbar and cloud-copy outputs of the mastery guide."""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document

from openwarframe.guide import (
    GuideFiles,
    build_mastery_plan,
    documents,
    find_browser,
    load_kb,
    render_mastery_html,
    write_guide_bundle,
    write_mastery_docx,
)
from openwarframe.guide.documents import copy_to_cloud, html_to_pdf, parse_targets
from test_guide import _report


@pytest.fixture(scope="module")
def plan():
    return build_mastery_plan(_report(), load_kb())


def test_word_document_has_every_section(plan, tmp_path: Path):
    path = write_mastery_docx(plan, tmp_path / "mastery-plan.docx")
    doc = Document(str(path))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Road to Mastery 30" in text
    for heading in (
        "Where you stand",
        "Do these first, tonight",
        "The path, one rank at a time",
        "The tests",
    ):
        assert heading in text
    assert "Operator Test" in text and "Vaykor Hek" in text
    assert len(doc.tables) >= 2  # rank table and budget table
    assert doc.tables[0].rows[1].cells[0].text == str(plan.ranks[0].rank)


def test_toolbar_links_real_files(plan):
    html = render_mastery_html(
        plan, GuideFiles(pdf="mastery-plan.pdf", docx="mastery-plan.docx", folder=r"C:\exports\x")
    )
    assert 'href="mastery-plan.pdf"' in html
    assert 'href="mastery-plan.docx" download' in html
    assert 'data-panel="cloud-gdrive"' in html and 'data-panel="cloud-onedrive"' in html
    assert "https://drive.google.com/drive/my-drive" in html and "https://onedrive.live.com/" in html
    assert 'data-copy="C:\\exports\\x"' in html
    assert '<button type="button" class="tool primary" data-print>' not in html
    assert "@media print" in html and ".tools, .tools-note, .cloud-panel, .chips{display:none" in html


def test_toolbar_falls_back_to_print_without_files(plan):
    html = render_mastery_html(plan)
    assert 'class="tool primary" data-print>' in html and "Save as PDF" in html
    assert ".docx" not in html.split("<nav")[1].split("</nav>")[0]
    assert 'data-copy="' not in html  # no folder path in the default/share page


def test_parse_targets_aliases_and_errors():
    assert parse_targets("OneDrive, google drive,gdrive") == ["onedrive", "gdrive"]
    assert parse_targets(None) == []
    with pytest.raises(ValueError):
        parse_targets("dropbox")


def test_copy_to_cloud_copies_into_openwarframe_folder(tmp_path: Path):
    src = tmp_path / "a.docx"
    src.write_bytes(b"x")
    used = copy_to_cloud([src], ["onedrive"], {"onedrive": tmp_path / "OneDrive"})
    assert (tmp_path / "OneDrive" / "OpenWarframe" / "a.docx").read_bytes() == b"x"
    assert used == {"onedrive": tmp_path / "OneDrive" / "OpenWarframe"}
    with pytest.raises(RuntimeError, match="Google Drive"):
        copy_to_cloud([src], ["gdrive"], {})


def test_bundle_writes_files_copies_and_warns(plan, tmp_path: Path):
    out = tmp_path / "out"
    bundle = write_guide_bundle(
        plan,
        out,
        pdf=False,
        copy_to=["onedrive", "gdrive"],
        folders={"onedrive": tmp_path / "OneDrive"},
    )
    assert bundle.html.exists() and bundle.docx and bundle.docx.exists() and bundle.pdf is None
    synced = tmp_path / "OneDrive" / "OpenWarframe"
    assert (synced / "mastery-plan.html").exists() and (synced / "mastery-plan.docx").exists()
    assert any("Google Drive" in w for w in bundle.warnings)
    html = bundle.html.read_text(encoding="utf-8")
    assert "Copied to OneDrive" in html
    assert 'class="tool primary" data-print>' in html  # pdf disabled, so the page offers Save as PDF


def test_share_mode_hides_local_folder(plan, tmp_path: Path):
    bundle = write_guide_bundle(plan, tmp_path, pdf=False, docx=False, share=True, folders={})
    html = bundle.html.read_text(encoding="utf-8")
    assert str(tmp_path) not in html


def test_find_browser_honours_env(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    fake = tmp_path / "chrome.exe"
    fake.write_bytes(b"")
    monkeypatch.setenv("OPENWARFRAME_BROWSER", str(fake))
    assert find_browser() == fake


def test_pdf_without_browser_explains(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    monkeypatch.setattr(documents, "find_browser", lambda: None)
    page = tmp_path / "p.html"
    page.write_text("<p>x</p>", encoding="utf-8")
    with pytest.raises(RuntimeError, match="OPENWARFRAME_BROWSER"):
        html_to_pdf(page, tmp_path / "p.pdf")


@pytest.mark.skipif(find_browser() is None, reason="no Chromium-based browser on this machine")
def test_pdf_prints_with_local_browser(plan, tmp_path: Path):
    bundle = write_guide_bundle(plan, tmp_path, docx=False, folders={})
    assert bundle.pdf is not None, bundle.warnings
    assert bundle.pdf.read_bytes()[:4] == b"%PDF"
    assert 'href="mastery-plan.pdf"' in bundle.html.read_text(encoding="utf-8")


def test_hidden_elements_stay_hidden(plan):
    html = render_mastery_html(plan)
    assert "[hidden]{display:none !important}" in html
    assert html.count('class="cloud-panel" id="cloud-') == 2
    assert html.count('" hidden>') >= 2
